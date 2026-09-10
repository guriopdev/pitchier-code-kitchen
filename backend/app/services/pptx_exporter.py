import io
from typing import Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from app.models.pitch import PitchDeck, Slide
from app.core.logger import logger


class PPTXExporterService:
    """Generates widescreen (16:9) PowerPoint presentations for PitchDeck payloads."""

    # Hallmark Modern Beige / Obsidian Precision Color Palette
    COLOR_BG = RGBColor(249, 247, 242)        # Warm Beige Linen (#F9F7F2)
    COLOR_CARD = RGBColor(255, 255, 255)      # Pure Card Surface (#FFFFFF)
    COLOR_INK = RGBColor(24, 24, 27)          # Charcoal Ink (#18181B)
    COLOR_SECONDARY = RGBColor(113, 113, 122) # Stone Grey (#71717A)
    COLOR_BORDER = RGBColor(228, 222, 207)    # Calibrated Hairline (#E4DECF)
    COLOR_ACCENT = RGBColor(194, 65, 12)      # Signal Terracotta (#C2410C)
    COLOR_EMERALD = RGBColor(21, 128, 61)     # Green (#15803D)
    COLOR_ROSE = RGBColor(190, 18, 60)        # Rose (#BE123C)
    COLOR_AMBER = RGBColor(180, 83, 9)        # Amber (#B45309)

    def export_deck_to_bytes(self, deck: PitchDeck) -> bytes:
        """Constructs a 16:9 PowerPoint file in-memory and returns the binary stream."""
        prs = Presentation()
        # Set 16:9 widescreen dimensions (13.333 x 7.5 inches)
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        blank_layout = prs.slide_layouts[6]

        # 1. Generate Title Cover Slide
        self._build_cover_slide(prs, deck, blank_layout)

        # 2. Generate 10 Content Slides
        for slide_data in deck.slides:
            self._build_content_slide(prs, slide_data, blank_layout)

        # Output to in-memory bytes buffer
        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)
        logger.info(f"Generated PPTX presentation for '{deck.company_name}' ({len(deck.slides) + 1} slides)")
        return buffer.getvalue()

    def _build_cover_slide(self, prs: Presentation, deck: PitchDeck, layout):
        slide = prs.slides.add_slide(layout)

        # Background fill
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = self.COLOR_BG
        bg.line.fill.background()

        # Company Title
        tx_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(11.0), Inches(1.5))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = deck.company_name
        p.font.size = Pt(48)
        p.font.bold = True
        p.font.color.rgb = self.COLOR_INK

        # One Liner Subtitle
        p2 = tf.add_paragraph()
        p2.text = deck.one_liner
        p2.font.size = Pt(20)
        p2.font.color.rgb = self.COLOR_SECONDARY
        p2.space_before = Pt(12)

        # Round & Ask Meta Bar
        meta_box = slide.shapes.add_textbox(Inches(1.2), Inches(5.2), Inches(11.0), Inches(1.0))
        tf_meta = meta_box.text_frame
        p_meta = tf_meta.paragraphs[0]
        p_meta.text = f"TARGET ROUND: {deck.target_round.upper()}   |   CAPITAL ASK: {deck.target_amount}"
        p_meta.font.size = Pt(14)
        p_meta.font.bold = True
        p_meta.font.color.rgb = self.COLOR_ACCENT

    def _build_content_slide(self, prs: Presentation, slide_data: Slide, layout):
        slide = prs.slides.add_slide(layout)

        # Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = self.COLOR_BG
        bg.line.fill.background()

        # Header bar: Slide Number & Classification
        header_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(0.6))
        tf_h = header_box.text_frame
        p_h = tf_h.paragraphs[0]
        p_h.text = f"{slide_data.slide_number:02d} / 10   ·   {slide_data.slide_type.value.upper().replace('_', ' ')}"
        p_h.font.size = Pt(11)
        p_h.font.bold = True
        p_h.font.color.rgb = self.COLOR_ACCENT

        # Headline & Subtitle
        title_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.1), Inches(11.3), Inches(1.4))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = slide_data.headline
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = self.COLOR_INK

        if slide_data.subtitle:
            p_sub = tf_t.add_paragraph()
            p_sub.text = slide_data.subtitle
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = self.COLOR_SECONDARY
            p_sub.space_before = Pt(4)

        # Left Column: Key Takeaways Card
        card_w = Inches(7.2)
        card_h = Inches(3.6)
        card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.6), card_w, card_h)
        card_shape.fill.solid()
        card_shape.fill.fore_color.rgb = self.COLOR_CARD
        card_shape.line.color.rgb = self.COLOR_BORDER

        tf_card = card_shape.text_frame
        tf_card.word_wrap = True
        tf_card.margin_left = Inches(0.4)
        tf_card.margin_right = Inches(0.4)
        tf_card.margin_top = Inches(0.4)

        for i, pt_text in enumerate(slide_data.key_points):
            p_pt = tf_card.paragraphs[0] if i == 0 else tf_card.add_paragraph()
            p_pt.text = f"•   {pt_text}"
            p_pt.font.size = Pt(13)
            p_pt.font.color.rgb = self.COLOR_INK
            if i > 0:
                p_pt.space_before = Pt(12)

        # Right Column: Metric Callout Cards
        if slide_data.metrics:
            metric_x = Inches(8.5)
            metric_w = Inches(3.8)
            metric_h = Inches(1.05)

            for idx, m in enumerate(slide_data.metrics[:3]):
                m_y = Inches(2.6) + idx * (metric_h + Inches(0.2))
                m_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, metric_x, m_y, metric_w, metric_h)
                m_card.fill.solid()
                m_card.fill.fore_color.rgb = self.COLOR_CARD
                m_card.line.color.rgb = self.COLOR_BORDER

                tf_m = m_card.text_frame
                tf_m.word_wrap = True
                tf_m.margin_left = Inches(0.3)
                tf_m.margin_top = Inches(0.15)

                p_lbl = tf_m.paragraphs[0]
                p_lbl.text = m.label.upper()
                p_lbl.font.size = Pt(9)
                p_lbl.font.color.rgb = self.COLOR_SECONDARY

                p_val = tf_m.add_paragraph()
                p_val.text = m.value
                p_val.font.size = Pt(20)
                p_val.font.bold = True
                p_val.font.color.rgb = self.COLOR_INK

                if m.context:
                    p_ctx = tf_m.add_paragraph()
                    p_ctx.text = m.context
                    p_ctx.font.size = Pt(9)
                    p_ctx.font.color.rgb = self.COLOR_SECONDARY

        # Speaker notes in PowerPoint notes section
        if slide_data.speaker_notes:
            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = f"FOUNDER SPOKEN SCRIPT:\n{slide_data.speaker_notes}\n\nINVESTOR CRITIQUE:\nStrengths: {', '.join(slide_data.investor_critique.strengths)}\nRed Flags: {', '.join(slide_data.investor_critique.red_flags)}"


pptx_exporter_service = PPTXExporterService()
